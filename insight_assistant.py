"""Insight Assistant — rule-based chat over insight_catalog.csv + rulebook.csv.
Floating overlay (bottom-right). Short answers, one watch-out line, chart + CSV per answer.
No LLM, no API key. Based on assistant_ui.py / retrieval.py (rules mode)."""
import difflib
import html
import inspect
import io
import re
from datetime import datetime

import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st

RED, GREEN, AMBER, DARK, LIGHT = "#E8002D", "#2A8050", "#906820", "#1C1C1C", "#FAFAF8"
REL = {"high": "High", "medium": "Medium"}
REL_RANK = {"High": 0, "Medium": 1, "Directional": 2}
METRIC_SYN = {
    "Attention": ["attention", "attentive", "stopping power", "noticed", "notice"],
    "Persuasion": ["persuasion", "persuasive", "purchase intent", "persuade"],
    "Likeability": ["likeability", "likability", "liking", "liked", "love"],
    "Experience Recall": ["experience recall", "recall", "memorable", "remember"],
    "Brand Linkage": ["brand linkage", "linkage", "branding", "brand link"],
    "Uniqueness": ["uniqueness", "unique", "distinctive"],
    "Shareability": ["shareability", "shareable", "share", "viral"],
    "Tiredness": ["tiredness", "tired", "tiring", "fatigue", "wear out", "wearout"],
    "SCD Score": ["scd", "scd score"],
}
LEVEL_WORDS = {"OU": r"\b(ous|operating units?|which ou|by ou|per ou)\b",
               "Market": r"\b(markets|countries|which market|which country|by market|per market)\b",
               "Brand": r"\b(brands|which brand|by brand|per brand)\b",
               "Category": r"\b(categories|which category|by category)\b"}
RX = {
    "help": r"^(help|\?|what can you do|how do i use|examples?)\b",
    "list": r"\b(list|show|which|what)\s+(all\s+)?(the\s+)?(features|metrics|markets|countries|brands|ous|scopes|categories)\b(?!\s+(work|drive|help|hurt|improve))",
    "rules": r"\b(alerts?|risks?|reversals?|revers\w*|flips?|conflicts?|trade.?offs?|opportunit\w*|underused|under.?used|anti.?patterns?|watch.?outs?|rulebook|outliers?|consensus|safe to brief)\b",
    "combo": r"\b(combinations?|combos?|together|pairs?|pairing)\b",
    "q_good": r"\b(best|help\w*|improv\w*|work\w*|good|strong\w*|top|drives?|driving|drivers?|safe)\b",
    "q_bad": r"\b(hurt\w*|worst|avoid\w*|harm\w*|bad|weak\w*|damag\w*|kills?)\b",
    "l_up": r"\b(increas\w*|rais\w*|boost\w*|lifts?|higher|positive\w*)\b",
    "l_down": r"\b(decreas\w*|reduc\w*|lower\w*|drops?|negative\w*)\b",
    "all_rel": r"\b(all findings|including (low|directional)|directional|everything)\b",
    "global": r"\b(global\w*|overall|whole dataset|all assets)\b",
    "mine": r"\b(my (scope|filters?|selection)|this (scope|selection)|here)\b",
    "follow": r"^(and|what about|how about|same for|now|also|and for)\b",
}
RULE_TYPES = {"heterogeneity": "Heterogeneity", "revers": "Heterogeneity", "flip": "Heterogeneity",
              "conflict": "Conflict", "trade": "Conflict", "opportunit": "Opportunity", "underused": "Opportunity",
              "under used": "Opportunity", "anti": "Anti-pattern", "outlier": "Outlier",
              "consensus": "Consensus", "safe to brief": "Consensus"}
STOP = set("a an the what which how do does is are of for in on to and with about show me tell insights insight "
           "findings finding please catalog evidence feature features affect affects impact on my scope filters "
           "selection where when why who it its be best worst top work works most least more less than vs versus "
           "drive drives help helps hurt hurts brand brands market markets ou ous globally global any all".split())


def _norm(v):
    return " ".join(re.findall(r"[a-z0-9&+]+", str(v).lower().replace("_", " ")))


def _e(v):
    return html.escape(str(v))


def _rerun():
    (getattr(st, "rerun", None) or getattr(st, "experimental_rerun"))()


def _kw(fn, **kw):
    params = inspect.signature(fn).parameters
    return {k: v for k, v in kw.items() if k in params}


# ── catalog vocabulary ────────────────────────────────────────────────────────
def _vocab(catalog, rulebook, feat_label):
    feats = {}
    for f, d in catalog[["feature", "feature_display"]].drop_duplicates().itertuples(index=False):
        if "+" in str(f): continue
        for name in {_norm(f), _norm(d), _norm(str(f).replace("_presence", "")), _norm(feat_label.get(f, ""))}:
            if len(name) >= 3: feats[name] = f
    metrics = sorted(catalog["metric_display"].dropna().unique())
    pairs = set(zip(catalog["filter"], catalog["filter_value"].astype(str)))
    if rulebook is not None:
        pairs |= set(zip(rulebook["scope"], rulebook["scope_value"].astype(str)))
    scopes = {_norm(v): (l, v) for l, v in pairs if l != "Global" and len(_norm(v)) >= 2}
    return feats, metrics, scopes


def _match_phrases(text, table):
    """Longest whole-word phrase matches; returns (payloads, text with matches removed)."""
    found = []
    for name in sorted(table, key=len, reverse=True):
        rx = r"(?<![a-z0-9])" + re.escape(name) + r"(?:e?s)?(?![a-z0-9])"
        if re.search(rx, text):
            found.append(table[name]); text = re.sub(rx, " ", text)
    return list(dict.fromkeys(found)), text


def parse(question, vocab, prev=None):
    feats_v, metrics_v, scopes_v = vocab
    raw = question.strip()
    t = _norm(raw)
    q = {"q": raw, "features": [], "metrics": [], "scopes": [], "level": None, "direction": None,
         "rules": [], "rules_any": False, "combo": False, "all_rel": False, "global": False, "mine": False,
         "n": 5, "help": False, "list": None}
    if re.search(RX["help"], t): q["help"] = True; return q
    m = re.search(RX["list"], t)
    if m and not re.search(r"\b(work|drive|help|hurt|improve|best|worst)\b", t):
        q["list"] = m.group(4); return q
    q["scopes"], t2 = _match_phrases(t, scopes_v)
    met_table = {s: mname for mname, syns in METRIC_SYN.items() if mname in metrics_v for s in syns}
    met_table.update({_norm(m_): m_ for m_ in metrics_v})
    q["metrics"], t2 = _match_phrases(t2, met_table)
    q["features"], t2 = _match_phrases(t2, feats_v)
    # typo tolerance on leftover words
    singles = {k: v for k, v in feats_v.items() if " " not in k}
    for tok in t2.split():
        if len(tok) >= 4 and tok not in STOP:
            hit = difflib.get_close_matches(tok, list(singles) + list(met_table), n=1, cutoff=.84)
            if hit:
                (q["features"] if hit[0] in singles else q["metrics"]).append(
                    singles.get(hit[0]) or met_table[hit[0]])
    q["features"] = list(dict.fromkeys(q["features"])); q["metrics"] = list(dict.fromkeys(q["metrics"]))
    q["missing_metrics"] = [m_ for m_, syns in METRIC_SYN.items() if m_ not in metrics_v
                            and any(re.search(r"(?<![a-z])" + re.escape(x) + r"(?![a-z])", t) for x in syns)]
    for lvl, rx in LEVEL_WORDS.items():
        if re.search(rx, t): q["level"] = lvl
    if re.search(r"\bwhere\b", t) and not q["scopes"]: q["level"] = q["level"] or "any"
    q["rules_any"] = bool(re.search(RX["rules"], t))
    q["rules"] = list(dict.fromkeys(v for k, v in RULE_TYPES.items() if k in t))
    q["combo"] = bool(re.search(RX["combo"], t))
    q["all_rel"] = bool(re.search(RX["all_rel"], t))
    q["global"] = bool(re.search(RX["global"], t))
    q["mine"] = bool(re.search(RX["mine"], t))
    n = re.search(r"\btop (\d+)\b", t)
    if n: q["n"] = max(1, min(int(n.group(1)), 10))
    t3 = re.sub(r"\btop \d+\b", " ", t)
    good = bool(re.search(RX["q_good"], t3)); bad = bool(re.search(RX["q_bad"], t3))
    up = bool(re.search(RX["l_up"], t3)); down = bool(re.search(RX["l_down"], t3))
    tired = q["metrics"] == ["Tiredness"]
    good |= (down if tired else up); bad |= (up if tired else down)
    q["direction"] = "good" if good and not bad else "bad" if bad and not good else None
    if q["rules_any"] and "avoid" in t and not q["rules"]: q["rules"] = ["Anti-pattern"]
    # follow-up: keep previous question, replace only what is new
    q["understood"] = any(q[k] for k in ("features", "metrics", "scopes", "level", "rules_any", "direction",
                                          "combo", "global", "mine", "missing_metrics"))
    only_scope = (q["scopes"] or q["level"] or q["global"] or q["mine"]) and not (
        q["features"] or q["metrics"] or q["rules_any"] or q["direction"])
    if prev and (re.search(RX["follow"], t) or only_scope):
        merged = dict(prev)
        for k in ("features", "metrics", "rules", "direction"):
            if q[k]: merged[k] = q[k]
        if q["scopes"] or q["level"] or q["global"] or q["mine"]:
            for k in ("scopes", "level", "global", "mine"): merged[k] = q[k]
        merged["q"] = raw; merged["n"] = q["n"] if n else prev.get("n", 5)
        return merged
    return q


# ── answering ────────────────────────────────────────────────────────────────
def _is_good(v, metric, lower_better_labels):
    return (v < 0) if metric in lower_better_labels else (v > 0)


def _scope_lbl(l, v):
    return "Global" if l == "Global" else f"{l} {v}"


def _scope_filter(df, q, page_scope, lcol, vcol):
    if q["scopes"]:
        m = pd.Series(False, index=df.index)
        for l, v in q["scopes"]: m |= (df[lcol] == l) & (df[vcol].astype(str) == str(v))
        return df[m], ", ".join(_scope_lbl(l, v) for l, v in q["scopes"])
    if q["level"] == "any": return df[df[lcol] != "Global"], "all scopes"
    if q["level"]: return df[df[lcol] == q["level"]], f"every {q['level']}"
    if (q["mine"] or page_scope["use"]) and page_scope["pair"] and not q["global"]:
        l, v = page_scope["pair"]
        return df[(df[lcol] == l) & (df[vcol].astype(str) == str(v))], _scope_lbl(l, v)
    return df[df[lcol] == "Global"], "Global"


def answer(question, catalog, rulebook, vocab, page_scope, lower_better_labels, rule_action, prev=None):
    q = parse(question, vocab, prev)
    out = {"q": q, "records": [], "kind": "text"}
    feats_v, metrics_v, scopes_v = vocab
    if q["help"]:
        out["html"] = ("<b>Ask me things like:</b><br>• What drives Attention in Fanta?<br>• Where does outdoors work best?"
                       "<br>• What should I avoid for Likeability?<br>• Any reversals for human presence?"
                       "<br>• what about LATAM? <span class='ia-dim'>(follow-up)</span>"
                       "<br><span class='ia-dim'>Say “my scope” to use your filters, “top 10” for more, “list markets” for names.</span>")
        return out
    if not q.get("understood", True) and not q["list"]:
        out["html"] = ("I couldn't match that to a feature, metric or market. "
                       "<span class='ia-dim'>Try “what drives Attention?”, “list features” or “help”.</span>")
        return out
    if q.get("missing_metrics") and not q["metrics"]:
        out["html"] = (f"<b>{_e(', '.join(q['missing_metrics']))}</b> isn't in the insight catalog — it covers "
                       f"{_e(', '.join(metrics_v))}. <span class='ia-dim'>Use the Feature Impact page for other metrics.</span>")
        return out
    if q["list"]:
        k = q["list"]
        if k == "features": vals = sorted({catalog.loc[catalog.feature == f, "feature_display"].iloc[0] for f in set(feats_v.values())})
        elif k == "metrics": vals = metrics_v
        else:
            lvl = {"markets": "Market", "countries": "Market", "brands": "Brand", "ous": "OU",
                   "categories": "Category", "scopes": None}[k]
            vals = sorted({v for l, v in scopes_v.values() if lvl is None or l == lvl})
        out["html"] = f"<b>{len(vals)} {k}:</b> " + _e(", ".join(map(str, vals[:40]))) + (" …" if len(vals) > 40 else "")
        return out

    fdisp = dict(zip(catalog["feature"], catalog["feature_display"]))
    # rulebook questions
    if q["rules_any"] and rulebook is not None and not q["direction"] == "good":
        rb = rulebook.copy()
        if q["rules"]: rb = rb[rb["rule_type"].isin(q["rules"])]
        if q["features"]: rb = rb[rb["feature"].isin(q["features"])]
        if q["metrics"]:
            rb = rb[rb["metric"].astype(str).apply(lambda m: m in ("", "All metrics") or any(x in m for x in q["metrics"])
                                                     or any(x.split()[0] in m for x in q["metrics"]))]
        rb_all = rb
        if q["scopes"] or q["level"] or q["global"] or ((q["mine"] or page_scope["use"]) and page_scope["pair"]):
            rb_s, where = _scope_filter(rb, q, page_scope, "scope", "scope_value")
            if where not in ("Global",) and not q["level"]:          # keep dataset-wide alerts alongside a named scope
                rb_s = pd.concat([rb_s, rb[rb["scope"] == "Global"]]); where += " + Global"
            rb = rb_s
        else:
            where = "all scopes"
        if rb.empty and where != "all scopes" and not q["scopes"]:
            rb, where = rb_all, f"all scopes (none for {where.replace(' + Global', '')})"
        rb = rb.sort_values("severity", key=lambda s: s.map({"high": 0, "medium": 1, "low": 2}))
        if rb.empty:
            what = (' / '.join(q['rules']) or 'alerts').lower()
            out["html"] = (f"No {what} found in any scope." if where.startswith("all scopes") else f"No {what} found for {_e(where)}.")
            return out
        rows = "".join(
            f"<div class='ia-row'><span class='ia-tag ia-{_e(r.rule_type[:4].lower())}'>{_e(r.rule_type)}</span>"
            f"<span class='ia-f'>{_e(fdisp.get(r.feature, str(r.feature).replace('_', ' ').title()))}</span>"
            f"<span class='ia-m'>{_e(_scope_lbl(r.scope, r.scope_value))}{' · HIGH' if r.severity == 'high' else ''}</span></div>"
            for r in rb.head(q["n"]).itertuples())
        types = rb.head(q["n"])["rule_type"].unique()
        tip = rule_action.get(types[0], "") if len(types) == 1 else ""
        out["html"] = (f"<b>{len(rb)} alert{'s' if len(rb) != 1 else ''} · {_e(where)}</b>"
                       + (f" <span class='ia-dim'>(showing {q['n']})</span>" if len(rb) > q["n"] else "")
                       + rows + (f"<div class='ia-dim'>→ {_e(tip)}</div>" if tip else ""))
        out["records"] = rb.drop(columns=[c for c in rb.columns if c.startswith("_")]).head(50).to_dict("records")
        out["kind"] = "rules"
        return out

    # catalog findings
    d = catalog.copy()
    if not q["combo"]: d = d[~d["insight_key"].astype(str).str.contains("Combination", case=False)]
    if q["features"]: d = d[d["feature"].isin(q["features"])]
    if q["metrics"]: d = d[d["metric_display"].isin(q["metrics"])]
    d_scoped, where = _scope_filter(d, q, page_scope, "filter", "filter_value")
    widened = ""
    if d_scoped.empty and not q["scopes"] and not q["level"] and where != "Global":
        d_scoped, widened, where = d[d["filter"] == "Global"], f"no findings for {where}, showing Global", "Global"
    if d_scoped.empty and where == "Global" and (q["features"] or q["metrics"]):
        d_scoped, where, widened = d[d["filter"] != "Global"], "all scopes", "no Global finding, so all scopes are shown"
    d = d_scoped
    d = d.assign(ia_u=pd.to_numeric(d["evidence_uplift_pp"], errors="coerce"),
                 ia_rel=d["confidence"].astype(str).str.lower().map(REL).fillna("Directional"))
    d = d[d["ia_u"].notna()]
    d["ia_good"] = [_is_good(u, m, lower_better_labels) for u, m in zip(d["ia_u"], d["metric_display"])]
    hidden = 0
    if not q["all_rel"]:
        strong = d[d["ia_rel"] != "Directional"]
        hidden = len(d) - len(strong)
        if len(strong): d = strong
    if q["direction"] == "good": d = d[d["ia_good"]]
    elif q["direction"] == "bad": d = d[~d["ia_good"]]
    if d.empty:
        what = " · ".join(filter(None, [", ".join(q["metrics"]), where]))
        out["html"] = (f"No tested findings for <b>{_e(what or 'that question')}</b>."
                       " <span class='ia-dim'>Try a feature or metric name, “list features”, or “help”.</span>")
        return out
    d = d.assign(ia_r=d["ia_rel"].map(REL_RANK), ia_a=d["ia_u"].abs()).sort_values(["ia_r", "ia_a"], ascending=[True, False])
    if not q["features"]: d = d.drop_duplicates(["feature", "metric_display", "filter", "filter_value"])
    show = d.head(q["n"])
    one_metric = show["metric_display"].nunique() == 1
    one_scope = show[["filter", "filter_value"]].drop_duplicates().shape[0] == 1
    # headline
    met = show["metric_display"].iloc[0] if one_metric else "your metrics"
    subj = (", ".join(catalog.loc[catalog.feature.isin(q["features"]), "feature_display"].drop_duplicates().head(2))
            if q["features"] else "Features")
    verb = {"good": f"improving {met}", "bad": f"worsening {met}"}.get(q["direction"], f"on {met}")
    sc = _scope_lbl(*show[["filter", "filter_value"]].iloc[0]) if one_scope else where
    head = f"<b>{_e(subj)} {_e(verb)} · {_e(sc)}</b>"
    if len(d) > len(show): head += f" <span class='ia-dim'>({len(show)} of {len(d)})</span>"
    rows = ""
    for r in show.itertuples():
        bits = ([] if one_metric else [r.metric_display]) + ([] if one_scope else [_scope_lbl(r.filter, r.filter_value)])
        rows += (f"<div class='ia-row'><span class='ia-f'>{_e(r.feature_display)}</span>"
                 f"<span class='ia-m'>{_e(' · '.join(bits))}</span>"
                 f"<span class='ia-v {'ia-good' if r.ia_good else 'ia-bad'}'>{r.ia_u:+.1f}pp</span>"
                 f"<span class='ia-rel ia-{r.ia_rel[:3].lower()}'>{r.ia_rel}</span>"
                 f"<span class='ia-n'>n={int(pd.to_numeric(r.evidence_n_has, errors='coerce') or 0):,}</span></div>")
    # one watch-out line
    warn = ""
    if rulebook is not None:
        rb = rulebook[rulebook["feature"].isin(show["feature"]) &
                      rulebook["rule_type"].isin(["Heterogeneity", "Conflict", "Anti-pattern"])]
        scopes_shown = set(zip(show["filter"], show["filter_value"].astype(str)))
        rb = rb[(rb["scope"] == "Global") | rb.apply(lambda x: (x["scope"], str(x["scope_value"])) in scopes_shown, axis=1)] if len(rb) else rb
        if len(rb):
            r0 = rb.sort_values("severity", key=lambda s: s.map({"high": 0, "medium": 1, "low": 2})).iloc[0]
            fd = catalog.loc[catalog.feature == r0["feature"], "feature_display"]
            more = f" (+{len(rb) - 1} more — ask “alerts for …”)" if len(rb) > 1 else ""
            warn = (f"<div class='ia-warn'>⚠ <b>{_e(r0['rule_type'])}</b> · {_e(fd.iloc[0] if len(fd) else r0['feature'])} · "
                    f"{_e(_scope_lbl(r0['scope'], r0['scope_value']))} — {_e(rule_action.get(r0['rule_type'], ''))}{_e(more)}</div>")
    notes = []
    if widened: notes.append(widened)
    if hidden and not q["all_rel"]: notes.append(f"{hidden} directional finding{'s' if hidden != 1 else ''} hidden")
    if any(m in lower_better_labels for m in show["metric_display"]): notes.append("Tiredness: lower is better")
    out["html"] = head + rows + warn + (f"<div class='ia-dim'>{_e(' · '.join(notes))}</div>" if notes else "")
    out["records"] = d.drop(columns=[c for c in d.columns if c.startswith("ia_")]).head(50).to_dict("records")
    out["chart_rows"] = show.head(10).assign(
        label=[f"{f}" + ("" if one_metric else f" · {m}") + ("" if one_scope else f" · {_scope_lbl(l, v)}")
               for f, m, l, v in zip(show["feature_display"], show["metric_display"], show["filter"], show["filter_value"])]
    )[["label", "ia_u", "ia_good", "ia_rel"]].to_dict("records")
    out["chart_title"] = f"{subj} {verb} · {sc}"
    out["bullets"] = [f"{r.feature_display}" + ("" if one_metric else f" · {r.metric_display}")
                      + ("" if one_scope else f" · {_scope_lbl(r.filter, r.filter_value)}")
                      + f": {r.ia_u:+.1f}pp ({r.ia_rel}, n={int(pd.to_numeric(r.evidence_n_has, errors='coerce') or 0):,})"
                      for r in show.itertuples()]
    out["warn"] = re.sub(r"<[^>]+>", "", warn).replace("&amp;", "&").strip()
    out["notes"] = " · ".join(notes)
    out["kind"] = "findings"
    return out


def chart_png(rows, title, dpi=200):
    fig = chart(rows, title)
    buf = io.BytesIO(); fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white"); plt.close(fig)
    return buf.getvalue()


def _plain(h):
    return html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", h or ""))).strip()


def chat_log(hist, deck):
    in_deck = {d["id"] for d in deck}
    return pd.DataFrame([{"time": t.get("time", ""), "question": t["q"], "answer": _plain(t["html"]),
                          "findings": len(t.get("records") or []), "in_deck": t.get("id") in in_deck} for t in hist])


def build_deck(deck, scope_txt=""):
    """PowerPoint (16:9): a title slide, then one slide per chart added to the deck."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    red, dark, grey = RGBColor(0xE8, 0x00, 0x2D), RGBColor(0x1C, 0x1C, 0x1C), RGBColor(0x6A, 0x66, 0x60)
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, lines, size=14, color=dark, bold=False):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
            p.space_after = Pt(6)
        return tf

    def bar(slide):
        r = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(.12)); r.fill.solid()
        r.fill.fore_color.rgb = red; r.line.fill.background()

    s0 = prs.slides.add_slide(blank); bar(s0)
    text(s0, .8, 2.4, 11.5, 1.2, ["Creative Intelligence — Insight findings"], 36, dark, True)
    text(s0, .8, 3.6, 11.5, 1, [f"{len(deck)} chart{'s' if len(deck) != 1 else ''} from the Insight Assistant · "
                                f"{datetime.now():%d %b %Y}" + (f" · {scope_txt}" if scope_txt else "")], 16, grey)
    for d in deck:
        sl = prs.slides.add_slide(blank); bar(sl)
        text(sl, .6, .35, 12.1, .9, [d["title"]], 24, dark, True)
        sl.shapes.add_picture(io.BytesIO(d.get("png") or chart_png(d["chart_rows"], "", dpi=220)), Inches(.6), Inches(1.35),
                              width=Inches(7.4))
        text(sl, 8.3, 1.4, 4.5, .8, ["Question: " + d["q"]], 12, grey)
        text(sl, 8.3, 2.2, 4.5, 3.6, ["• " + b for b in d.get("bullets", [])], 13)
        if d.get("warn"): text(sl, 8.3, 5.6, 4.5, 1, [d["warn"]], 11, RGBColor(0x90, 0x68, 0x20))
        text(sl, .6, 6.85, 12.1, .4, ["Source: insight_catalog.csv / rulebook.csv · Uplift = assets with the feature minus "
                                      "assets without it (pp) · Reliability: High = ** or *** with 50+ assets each side"
                                      + (" · " + d["notes"] if d.get("notes") else "")], 9, grey)
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()


def chart(rows, title):
    d = pd.DataFrame(rows).iloc[::-1]
    fig, ax = plt.subplots(figsize=(4.6, max(1.8, .42 * len(d) + .8)))
    fig.patch.set_facecolor("white"); ax.set_facecolor("white")
    bars = ax.barh(d["label"].str.slice(0, 34), d["ia_u"], color=[GREEN if g else RED for g in d["ia_good"]], height=.6)
    for b, r in zip(bars, d["ia_rel"]): b.set_alpha(.95 if r == "High" else .7 if r == "Medium" else .4)
    ax.axvline(0, color="#999", lw=.8)
    for i, (v, r) in enumerate(zip(d["ia_u"], d["ia_rel"])):
        ax.text(v, i, f" {v:+.1f}pp " + ("" if r == "High" else f"({r})"), va="center",
                ha="left" if v >= 0 else "right", fontsize=7.5, color=DARK)
    lo, hi = min(0, d["ia_u"].min()), max(0, d["ia_u"].max()); pad = (hi - lo) * .35 or 1
    ax.set_xlim(lo - (pad if lo < 0 else 0), hi + (pad if hi > 0 else 0))
    ax.set_xlabel("Uplift (pp)", fontsize=8); ax.tick_params(labelsize=8)
    ax.set_title(title[:70], fontsize=8.5, loc="left", color=DARK)
    for s in ("top", "right", "left"): ax.spines[s].set_visible(False)
    plt.tight_layout()
    return fig


# ── UI ───────────────────────────────────────────────────────────────────────
CSS = """<style>
div[data-testid="stVerticalBlock"]:has([data-ia-launch]):not(:has(div[data-testid="stVerticalBlock"] [data-ia-launch])){
 position:fixed!important;right:26px!important;bottom:24px!important;width:auto!important;z-index:100002!important;gap:0!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-launch]):not(:has(div[data-testid="stVerticalBlock"] [data-ia-launch])) button{
 background:linear-gradient(135deg,#E8002D,#B5001F)!important;color:#fff!important;border:0!important;border-radius:28px!important;
 padding:.62rem 1.25rem!important;box-shadow:0 8px 26px rgba(232,0,45,.35)!important;font-weight:600!important;width:auto!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-launch]):not(:has(div[data-testid="stVerticalBlock"] [data-ia-launch])) button p{color:#fff!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-panel]):not(:has(div[data-testid="stVerticalBlock"] [data-ia-panel])){
 position:fixed!important;right:26px!important;bottom:84px!important;width:min(420px,calc(100vw - 32px))!important;
 max-height:min(690px,calc(100vh - 120px))!important;overflow-y:auto!important;overscroll-behavior:contain!important;
 z-index:100001!important;background:#fff!important;border:1px solid #EAE8E2!important;border-radius:18px!important;
 padding:0 16px 14px 16px!important;box-shadow:0 18px 60px rgba(0,0,0,.22)!important;gap:.45rem!important;}
div[data-testid="stElementContainer"]:has([data-ia-launch]),div[data-testid="stElementContainer"]:has([data-ia-panel]),
div.element-container:has([data-ia-launch]),div.element-container:has([data-ia-panel]){display:none!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-panel]) .stButton>button{border-radius:16px!important;font-size:.76rem!important;
 padding:.18rem .7rem!important;min-height:0!important;width:auto!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-panel]) .stDownloadButton>button{border-radius:16px!important;font-size:.76rem!important;
 padding:.18rem .7rem!important;min-height:0!important;background:#fff!important;border:1px solid #DAD6D0!important;color:#6A6660!important;}
div[data-testid="stVerticalBlock"]:has([data-ia-panel]) [data-testid="stForm"]{border:1px solid #EAE8E2!important;border-radius:14px!important;padding:.5rem .7rem!important;}
.ia-head{display:flex;align-items:center;gap:.7rem;margin:0 -16px .2rem -16px;padding:.9rem 16px;
 background:linear-gradient(135deg,#E8002D,#B5001F);border-radius:18px 18px 0 0;color:#fff;}
.ia-av{width:34px;height:34px;border-radius:50%;background:rgba(255,255,255,.18);display:flex;align-items:center;
 justify-content:center;font-weight:700;font-size:.8rem;border:1px solid rgba(255,255,255,.35);}
.ia-title{font-family:'Merriweather',Georgia,serif;font-size:1rem;line-height:1.1;}
.ia-sub{font-size:.7rem;opacity:.85;margin-top:.15rem;} .ia-dot{display:inline-block;width:7px;height:7px;border-radius:50%;background:#7CF0A8;margin-right:.3rem;}
.ia-ctx{font-size:.7rem;color:#888;margin:.1rem 0 .1rem 0;}
.ia-u{display:flex;justify-content:flex-end;margin:.35rem 0 .1rem 0;}
.ia-u div{background:#E8002D;color:#fff;border-radius:16px 16px 4px 16px;padding:.45rem .8rem;font-size:.83rem;max-width:85%;line-height:1.45;}
.ia-a{display:flex;gap:.45rem;align-items:flex-start;margin:.15rem 0 .1rem 0;}
.ia-a .ia-av{background:#FFF0F0;color:#E8002D;border:1px solid #F5C8D0;width:26px;height:26px;font-size:.62rem;flex-shrink:0;}
.ia-b{background:#F7F6F3;border-radius:4px 16px 16px 16px;padding:.6rem .8rem;font-size:.82rem;color:#1C1C1C;line-height:1.5;flex:1;}
.ia-row{display:flex;align-items:center;gap:.4rem;padding:.28rem 0;border-bottom:1px solid #ECEAE5;font-size:.8rem;}
.ia-row:last-of-type{border-bottom:0;}
.ia-f{font-weight:600;flex:1;min-width:0;} .ia-m{color:#888;font-size:.72rem;}
.ia-v{font-weight:700;white-space:nowrap;} .ia-good{color:#2A8050;} .ia-bad{color:#C00020;}
.ia-rel{font-size:.62rem;font-weight:700;border-radius:10px;padding:.05rem .45rem;}
.ia-hig{background:#EAFAF1;color:#27AE60;} .ia-med{background:#FEF9E7;color:#E67E22;} .ia-dir{background:#FFF7E6;color:#906820;}
.ia-n{color:#AAA;font-size:.68rem;white-space:nowrap;}
.ia-tag{font-size:.62rem;font-weight:700;text-transform:uppercase;letter-spacing:.04em;border-radius:10px;padding:.08rem .45rem;background:#F2F2F2;color:#555;}
.ia-hete{background:#FFFBF0;color:#906820;} .ia-conf,.ia-anti{background:#FFF0F0;color:#C00020;}
.ia-oppo,.ia-cons{background:#F0FFF4;color:#1A7040;} .ia-outl{background:#EBF3FB;color:#1A3080;} .ia-boun{background:#F8F0FF;color:#7030A0;}
.ia-warn{margin-top:.45rem;background:#FFF8E6;border:1px solid #F0D890;border-radius:8px;padding:.35rem .55rem;font-size:.75rem;color:#5A4210;}
.ia-dim{color:#999;font-size:.72rem;margin-top:.3rem;}
.ia-empty{font-size:.8rem;color:#777;padding:.4rem 0 .2rem 0;}
</style>"""

SUGGESTIONS = ["What drives Attention here?", "What should I avoid?", "Any reversals I should know about?"]


def render_assistant(catalog, rulebook, page_scope, lower_better_labels, rule_action, feat_label):
    """page_scope: {"use": bool default, "pair": (level, value) or None, "note": str}."""
    st.markdown(CSS, unsafe_allow_html=True)
    ss = st.session_state
    with st.container():
        st.markdown('<span data-ia-launch="1"></span>', unsafe_allow_html=True)
        if st.button("✕  Close" if ss.get("ia_open") else "✦  Ask insights", key="ia_toggle"):
            ss["ia_open"] = not ss.get("ia_open", False); _rerun()
    if not ss.get("ia_open"):
        return
    with st.container():
        st.markdown('<span data-ia-panel="1"></span>', unsafe_allow_html=True)
        st.markdown('<div class="ia-head"><div class="ia-av">AI</div><div><div class="ia-title">Insight Assistant</div>'
                    '<div class="ia-sub"><span class="ia-dot"></span>Answers from the Insight Catalog &amp; Rulebook</div></div></div>',
                    unsafe_allow_html=True)
        if catalog is None or catalog.empty:
            st.markdown('<div class="ia-empty">I can\'t find <b>insight_catalog.csv</b>. Put it in the same folder as '
                        'creative_intelligence.py — it reloads automatically.</div>', unsafe_allow_html=True)
            return
        vocab = _vocab(catalog, rulebook, feat_label)
        hist = ss.setdefault("ia_hist", [])

        c1, c2 = st.columns([3, 1.3])
        with c1:
            use = st.checkbox("Use my filters", value=True, key="ia_use",
                              help="Answer for the OU / Brand / Market selected in the page filters.")
        with c2:
            if hist and st.button("↺ New chat", key="ia_new"):
                ss["ia_hist"] = []; ss.pop("ia_prev", None); _rerun()   # the deck is kept across new chats
        ps = dict(page_scope, use=use)
        ctx = (f"Scope: {_scope_lbl(*ps['pair'])} (from your filters)" if use and ps["pair"] else "Scope: Global unless you name one")
        st.markdown(f'<div class="ia-ctx">{_e(ctx)}{" · " + _e(ps["note"]) if use and ps.get("note") else ""}</div>',
                    unsafe_allow_html=True)

        pending = ss.pop("ia_pending", None)
        if not hist and not pending:
            st.markdown('<div class="ia-a"><div class="ia-av">AI</div><div class="ia-b">Hi! Ask which creative features '
                        'drive a metric, what to avoid, or where findings reverse. Type <b>help</b> for examples.</div></div>',
                        unsafe_allow_html=True)
            for i, sug in enumerate(SUGGESTIONS):
                if st.button(sug, key=f"ia_sug_{i}"):
                    ss["ia_pending"] = sug; _rerun()

        if len(hist) > 1:
            with st.expander(f"Earlier ({len(hist) - 1})"):
                for t in hist[:-1]:
                    st.markdown(f'<div class="ia-u"><div>{_e(t["q"])}</div></div>'
                                f'<div class="ia-a"><div class="ia-av">AI</div><div class="ia-b">{t["html"]}</div></div>',
                                unsafe_allow_html=True)
        if hist:
            i, t = len(hist) - 1, hist[-1]
            st.markdown(f'<div class="ia-u"><div>{_e(t["q"])}</div></div>'
                        f'<div class="ia-a"><div class="ia-av">AI</div><div class="ia-b">{t["html"]}</div></div>',
                        unsafe_allow_html=True)
            if t.get("records"):
                b1, b2, _ = st.columns([1, 1, 1.4])
                with b1:
                    if t.get("chart_rows") and st.button("▾ Hide chart" if t.get("show_chart") else "📊 Chart", key=f"ia_ch_{i}"):
                        t["show_chart"] = not t.get("show_chart"); _rerun()
                with b2:
                    st.download_button("⬇ CSV", pd.DataFrame(t["records"]).to_csv(index=False).encode("utf-8"),
                                       file_name="assistant_findings.csv", mime="text/csv", key=f"ia_dl_{i}")
                if t.get("show_chart") and t.get("chart_rows"):
                    png = chart_png(t["chart_rows"], t.get("chart_title", ""))
                    st.image(png)
                    deck = ss.setdefault("ia_deck", [])
                    added = any(d["id"] == t["id"] for d in deck)
                    c1, c2, _ = st.columns([1, 1.25, 1])
                    with c1:
                        st.download_button("⬇ PNG", png, file_name="assistant_chart.png", mime="image/png", key=f"ia_png_{i}")
                    with c2:
                        if st.button("✓ In deck" if added else "＋ Add to deck", key=f"ia_add_{i}"):
                            ss["ia_deck"] = ([d for d in deck if d["id"] != t["id"]] if added else deck + [{
                                "id": t["id"], "q": t["q"], "title": t.get("chart_title") or t["q"],
                                "chart_rows": t["chart_rows"], "bullets": t.get("bullets", []),
                                "png": chart_png(t["chart_rows"], "", dpi=220),
                                "warn": t.get("warn", ""), "notes": t.get("notes", "")}])
                            _rerun()

        # session tracking: chat log + deck of charts
        deck = ss.get("ia_deck", [])
        if hist:
            d1, d2, _ = st.columns([1.15, 1.25, .6])
            with d1:
                st.download_button("⬇ Chat log", chat_log(hist, deck).to_csv(index=False).encode("utf-8"),
                                   file_name=f"assistant_chat_{datetime.now():%Y%m%d_%H%M}.csv", mime="text/csv", key="ia_log")
            with d2:
                if deck:
                    try:
                        pptx = build_deck(deck, ctx)
                        st.download_button(f"📑 Deck ({len(deck)})", pptx,
                                           file_name=f"insight_charts_{datetime.now():%Y%m%d_%H%M}.pptx",
                                           mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                                           key="ia_deck_dl")
                    except ImportError:
                        st.markdown('<div class="ia-dim">Deck export needs <b>python-pptx</b>: pip install python-pptx</div>',
                                    unsafe_allow_html=True)
                else:
                    st.markdown('<div class="ia-dim">Show a chart and press ＋ Add to deck to build a PowerPoint.</div>',
                                unsafe_allow_html=True)

        with st.form("ia_form", clear_on_submit=True):
            f1, f2 = st.columns([4, 1])
            with f1:
                question = st.text_input("Question", placeholder="Ask about a feature, metric or market…",
                                         **_kw(st.text_input, label_visibility="collapsed"))
            with f2:
                sent = st.form_submit_button("Send ➤")
        question = pending or (question.strip() if sent else "")
        if question:
            try:
                res = answer(question[:500], catalog, rulebook, vocab, ps, lower_better_labels, rule_action, ss.get("ia_prev"))
            except Exception as exc:   # never break the page because of the assistant
                res = {"html": f"Sorry — I couldn't process that ({_e(type(exc).__name__)}).", "q": None}
            if res.get("q") and not res["q"].get("help") and not res["q"].get("list"):
                ss["ia_prev"] = res["q"]
            ss["ia_seq"] = ss.get("ia_seq", 0) + 1
            hist.append({"id": ss["ia_seq"], "time": f"{datetime.now():%Y-%m-%d %H:%M}",
                         "bullets": res.get("bullets", []), "warn": res.get("warn", ""), "notes": res.get("notes", ""),
                         "q": question, "html": res["html"], "records": res.get("records", []),
                         "chart_rows": res.get("chart_rows"), "chart_title": res.get("chart_title", ""),
                         "show_chart": False})
            ss["ia_hist"] = hist[-30:]
            _rerun()
